"""Serfis PowerPoint template, built from the purple Serfis design-system artifact."""
import copy
from pathlib import Path
import re
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import PP_PLACEHOLDER, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent          # brand/build
BG = str(HERE / "bg")
LOGO = str(HERE.parent / "assets" / "serfis-logo.png")
OUT = sys.argv[1] if len(sys.argv) > 1 else str(HERE.parent / "Serfis_Template.pptx")

# ---- brand tokens (from the artifact's :root) ----
C = dict(bg="06030F", bg2="0F032D", story="07031A", violet="4E2BCC", violet2="6A3FE6", lavender="905BF4",
         heliotrope="C084FC", blue="3B5BFF", blue_deep="2438C8", ink="0F032D", text="EFEFEF",
         head="F6F1FF",                 # slide headline white
         text2="B7B6B9",                # body copy: text at ~.76 over the ground
         text3="7B797F",                # labels: text at .5
         muted="A9A8AC",                # text at .7
         line="302D37", glass="10062C", creative="D946EF", green="3BE08C", rec="FF3D6E", paper="FFFFFF",
         on_paper="6D3FE0")
DISPLAY, BODY, MONO, WORD = "Unbounded", "DM Sans", "DM Mono", "Fredoka"
W_IN, H_IN = 13.333, 7.5
LM = 0.85                                  # left margin
CW = W_IN - 2 * LM                         # content width


def E(inches):
    return int(round(inches * 914400))


# =====================================================================================
# XML helpers
# =====================================================================================
def X(xml):
    return parse_xml(xml.replace("<p:", f"<p:", 1) if False else _wrap(xml))


def _wrap(xml):
    # inject namespace declarations into the first tag
    m = re.match(r"<([\w:]+)", xml)
    tag = m.group(0)
    return xml.replace(tag, tag + " " + nsdecls("a", "p", "r"), 1)


def solid(hex_, alpha=None):
    a = f'<a:alpha val="{int(alpha * 1000)}"/>' if alpha is not None else ""
    return f'<a:solidFill><a:srgbClr val="{hex_}">{a}</a:srgbClr></a:solidFill>'


def grad(stops, ang):
    gs = "".join(f'<a:gs pos="{int(p * 1000)}"><a:srgbClr val="{c}"/></a:gs>' for p, c in stops)
    return f'<a:gradFill rotWithShape="1"><a:gsLst>{gs}</a:gsLst><a:lin ang="{ang}" scaled="0"/></a:gradFill>'


BTN_GRAD = grad([(0, "7F52FF"), (50, "6A3FE6"), (100, "3F46E8")], 2700000)
TAG_GRAD = grad([(0, C["heliotrope"]), (100, C["blue"])], 0)


def run_xml(text, size, color, font, bold=False, spc=0, cap=False, italic=False):
    text = text.replace("&", "&amp;").replace("<", "&lt;")
    return (f'<a:r><a:rPr lang="en-US" sz="{int(size * 100)}" b="{int(bold)}" i="{int(italic)}" spc="{spc}"'
            f'{" cap=" + chr(34) + "all" + chr(34) if cap else ""} dirty="0">{solid(color)}'
            f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/></a:rPr><a:t>{text}</a:t></a:r>')


def shape_xml(sid, name, x, y, w, h, geom="rect", fill="", line="", text_xml="", adj=None, anchor="ctr",
              insets=(0, 0, 0, 0), txbox=False, algn="l", custgeom=None):
    av = f'<a:avLst><a:gd name="adj" fmla="val {adj}"/></a:avLst>' if adj is not None else "<a:avLst/>"
    geo = custgeom or f'<a:prstGeom prst="{geom}">{av}</a:prstGeom>'
    fill = fill or "<a:noFill/>"
    line = line or "<a:ln><a:noFill/></a:ln>"
    body = ""
    if text_xml:
        l, t, r, b = (E(v) for v in insets)
        body = (f'<p:txBody><a:bodyPr wrap="square" lIns="{l}" tIns="{t}" rIns="{r}" bIns="{b}" anchor="{anchor}">'
                f'<a:noAutofit/></a:bodyPr><a:lstStyle/><a:p><a:pPr algn="{algn}"/>{text_xml}</a:p></p:txBody>')
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr{" txBox=" + chr(34) + "1" + chr(34) if txbox else ""}/>'
            f'<p:nvPr userDrawn="1"/></p:nvSpPr><p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(h)}"/></a:xfrm>'
            f'{geo}{fill}{line}</p:spPr>{body}</p:sp>')


def corner_xml(sid, name, x, y, s, which):
    """Viewfinder corner bracket: an open L, heliotrope at 60%."""
    n = E(s)
    pts = {"tl": [(0, n), (0, 0), (n, 0)], "tr": [(0, 0), (n, 0), (n, n)],
           "bl": [(0, 0), (0, n), (n, n)], "br": [(n, 0), (n, n), (0, n)]}[which]
    path = f'<a:moveTo><a:pt x="{pts[0][0]}" y="{pts[0][1]}"/></a:moveTo>' + "".join(
        f'<a:lnTo><a:pt x="{px}" y="{py}"/></a:lnTo>' for px, py in pts[1:])
    cg = (f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/><a:pathLst>'
          f'<a:path w="{n}" h="{n}" fill="none">{path}</a:path></a:pathLst></a:custGeom>')
    ln = f'<a:ln w="12700" cap="flat">{solid(C["heliotrope"], 60)}<a:miter lim="800000"/></a:ln>'
    return shape_xml(sid, name, x, y, s, s, custgeom=cg, line=ln)


def notch_geom(w, h, n=0.22):
    W_, H_, N = E(w), E(h), E(n)
    return (f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/><a:pathLst>'
            f'<a:path w="{W_}" h="{H_}"><a:moveTo><a:pt x="0" y="0"/></a:moveTo><a:lnTo><a:pt x="{W_}" y="0"/></a:lnTo>'
            f'<a:lnTo><a:pt x="{W_}" y="{H_ - N}"/></a:lnTo><a:lnTo><a:pt x="{W_ - N}" y="{H_}"/></a:lnTo>'
            f'<a:lnTo><a:pt x="0" y="{H_}"/></a:lnTo><a:close/></a:path></a:pathLst></a:custGeom>')


def pic_xml(sid, name, rid, x, y, w, h):
    return (f'<p:pic><p:nvPicPr><p:cNvPr id="{sid}" name="{name}" descr="Serfis logo"/><p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr>'
            f'<p:nvPr userDrawn="1"/></p:nvPicPr><p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
            f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(h)}"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>')


def placeholder_xml(sid, name, idx, x, y, w, h, lvl1_rpr, sample, ph_type="body", anchor="ctr"):
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>'
            f'<p:nvPr><p:ph type="{ph_type}" sz="quarter" idx="{idx}"/></p:nvPr></p:nvSpPr>'
            f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(w)}" cy="{E(h)}"/></a:xfrm></p:spPr>'
            f'<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}"><a:noAutofit/></a:bodyPr>'
            f'<a:lstStyle><a:lvl1pPr marL="0" indent="0"><a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>{lvl1_rpr}</a:lvl1pPr></a:lstStyle>'
            f'<a:p><a:r><a:rPr lang="en-US"/><a:t>{sample}</a:t></a:r></a:p></p:txBody></p:sp>')


def rpr(size, color, font, bold=False, spc=0, cap=False):
    return (f'<a:defRPr sz="{int(size * 100)}" b="{int(bold)}" spc="{spc}"{" cap=" + chr(34) + "all" + chr(34) if cap else ""}>'
            f'{solid(color)}<a:latin typeface="{font}"/><a:cs typeface="{font}"/></a:defRPr>')


# =====================================================================================
# Presentation, theme, master
# =====================================================================================
prs = Presentation()
prs.slide_width, prs.slide_height = E(W_IN), E(H_IN)
master = prs.slide_master

# theme: dark scheme + brand fonts
theme_part = master.part.part_related_by(RT.THEME)
theme = theme_part.blob.decode("utf8")
clr = ('<a:clrScheme name="Serfis"><a:dk1><a:srgbClr val="06030F"/></a:dk1><a:lt1><a:srgbClr val="EFEFEF"/></a:lt1>'
       '<a:dk2><a:srgbClr val="0F032D"/></a:dk2><a:lt2><a:srgbClr val="C084FC"/></a:lt2>'
       '<a:accent1><a:srgbClr val="6A3FE6"/></a:accent1><a:accent2><a:srgbClr val="905BF4"/></a:accent2>'
       '<a:accent3><a:srgbClr val="C084FC"/></a:accent3><a:accent4><a:srgbClr val="3B5BFF"/></a:accent4>'
       '<a:accent5><a:srgbClr val="D946EF"/></a:accent5><a:accent6><a:srgbClr val="3BE08C"/></a:accent6>'
       '<a:hlink><a:srgbClr val="C084FC"/></a:hlink><a:folHlink><a:srgbClr val="905BF4"/></a:folHlink></a:clrScheme>')
theme = re.sub(r"<a:clrScheme .*?</a:clrScheme>", clr, theme, flags=re.S)
theme = re.sub(r'(<a:majorFont>\s*<a:latin typeface=")[^"]*', rf"\g<1>{DISPLAY}", theme)
theme = re.sub(r'(<a:minorFont>\s*<a:latin typeface=")[^"]*', rf"\g<1>{BODY}", theme)
theme = re.sub(r'<a:theme ([^>]*)name="[^"]*"', r'<a:theme \1name="Serfis"', theme, count=1)
theme_part._blob = theme.encode("utf8")

# master: map scheme so bg1=dark, tx1=light
cm = master._element.find(qn("p:clrMap"))
for k, v in dict(bg1="dk1", tx1="lt1", bg2="dk2", tx2="lt2").items():
    cm.set(k, v)


def lvl(n, size, mar, indent, bullet, color=C["text2"], spc=0):
    return (f'<a:lvl{n}pPr marL="{mar}" indent="{indent}" algn="l" defTabSz="914400" rtl="0" eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
            f'<a:lnSpc><a:spcPct val="112000"/></a:lnSpc><a:spcBef><a:spcPts val="{spc}"/></a:spcBef>'
            f'<a:buClr><a:srgbClr val="{C["heliotrope"]}"/></a:buClr><a:buFont typeface="Arial"/><a:buChar char="{bullet}"/>'
            f'<a:defRPr sz="{size}" kern="1200">{solid(color)}<a:latin typeface="+mn-lt"/><a:ea typeface="+mn-ea"/><a:cs typeface="+mn-cs"/></a:defRPr></a:lvl{n}pPr>')


body_style = "".join([lvl(1, 1800, 285750, -285750, "•", spc=900), lvl(2, 1600, 628650, -228600, "–", spc=500),
                      lvl(3, 1400, 971550, -228600, "•", spc=400), lvl(4, 1400, 1314450, -228600, "–", spc=300),
                      lvl(5, 1400, 1657350, -228600, "•", spc=300)])
tx = master._element.find(qn("p:txStyles"))
title_style = (f'<p:titleStyle {nsdecls("a", "p", "r")}><a:lvl1pPr algn="l" defTabSz="914400" rtl="0" eaLnBrk="1" latinLnBrk="0" hangingPunct="1">'
               f'<a:lnSpc><a:spcPct val="100000"/></a:lnSpc><a:spcBef><a:spcPct val="0"/></a:spcBef><a:buNone/>'
               f'<a:defRPr sz="3000" b="1" cap="all" spc="-30" kern="1200">{solid(C["head"])}<a:latin typeface="+mj-lt"/><a:ea typeface="+mj-ea"/><a:cs typeface="+mj-cs"/></a:defRPr>'
               f'</a:lvl1pPr></p:titleStyle>')
tx.replace(tx.find(qn("p:titleStyle")), parse_xml(title_style))
tx.replace(tx.find(qn("p:bodyStyle")), parse_xml(f'<p:bodyStyle {nsdecls("a", "p", "r")}>{body_style}</p:bodyStyle>'))

# ---- layout bookkeeping ----
layouts = {l.name: l for l in prs.slide_layouts}
L_TITLE = layouts["Title Slide"]
L_CONTENT = layouts["Title and Content"]
L_SECTION = layouts["Section Header"]
L_TWO = layouts["Two Content"]
L_COMPARE = layouts["Comparison"]
L_TITLEONLY = layouts["Title Only"]
L_BLANK = layouts["Blank"]
L_CAPTION = layouts["Content with Caption"]
L_PICTURE = layouts["Picture with Caption"]
L_LIGHT = layouts["Title and Vertical Text"]
L_BAND = layouts["Vertical Title and Text"]


def ph_by_idx(layout, idx):
    for p in layout.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)


def place(ph, x, y, w, h):
    ph.left, ph.top, ph.width, ph.height = E(x), E(y), E(w), E(h)


def style_ph(ph, size, color, font, bold=False, anchor="t", cap=None, bullets=None, spc=0, algn="l", lnspc=None):
    """Write the placeholder's own list style so the layout (not each slide) owns the look."""
    txBody = ph._element.txBody
    bp = txBody.find(qn("a:bodyPr"))
    for k in ("vert", "rot"):
        if k in bp.attrib:
            del bp.attrib[k]
    for k, v in (("anchor", anchor), ("lIns", "0"), ("tIns", "0"), ("rIns", "0"), ("bIns", "0"), ("wrap", "square")):
        bp.set(k, v)
    for ch in list(bp):
        bp.remove(ch)
    bp.append(parse_xml(f'<a:noAutofit {nsdecls("a")}/>'))
    lst = txBody.find(qn("a:lstStyle"))
    for ch in list(lst):
        lst.remove(ch)
    bu = "<a:buNone/>" if bullets is None else (f'<a:buClr><a:srgbClr val="{C["heliotrope"]}"/></a:buClr>'
                                                f'<a:buFont typeface="Arial"/><a:buChar char="{bullets}"/>')
    mar = 'marL="0" indent="0"' if bullets is None else 'marL="285750" indent="-285750"'
    ls = f'<a:lnSpc><a:spcPct val="{lnspc}"/></a:lnSpc>' if lnspc else ""
    capattr = ' cap="all"' if cap else (' cap="none"' if cap is False else "")
    lst.append(parse_xml(
        f'<a:lvl1pPr {nsdecls("a")} {mar} algn="{algn}">{ls}<a:spcBef><a:spcPts val="0"/></a:spcBef>{bu}'
        f'<a:defRPr sz="{int(size * 100)}" b="{int(bold)}" spc="{spc}"{capattr}>{solid(color)}'
        f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/></a:defRPr></a:lvl1pPr>'))


def set_background(layout, image=None, color=None):
    cSld = layout._element.find(qn("p:cSld"))
    old = cSld.find(qn("p:bg"))
    if old is not None:
        cSld.remove(old)
    if image:
        _, rid = layout.part.get_or_add_image_part(image)
        bg = (f'<p:bg {nsdecls("a", "p", "r")}><p:bgPr><a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/>'
              f'<a:srcRect/><a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
    else:
        bg = f'<p:bg {nsdecls("a", "p", "r")}><p:bgPr>{solid(color)}<a:effectLst/></p:bgPr></p:bg>'
    cSld.insert(0, parse_xml(bg))


def decorate(layout, xml_list):
    """Static art goes at the back of the layout (index 2 = after the group properties)."""
    tree = layout.shapes._spTree
    for i, xml in enumerate(xml_list):
        tree.insert(2 + i, X(xml))


def add_placeholder(layout, xml):
    layout.shapes._spTree.append(X(xml))


def next_ids(base):
    n = [base]

    def f():
        n[0] += 1
        return n[0]
    return f


def logo_rid(layout):
    return layout.part.get_or_add_image_part(LOGO)[1]


def strip_orient(layout):
    for sp in layout.shapes._spTree.iter(qn("p:ph")):
        if "orient" in sp.attrib:
            del sp.attrib["orient"]


def finish_footers(layout, ink=C["text3"], show=True):
    """Footer text left, slide number right; mono 9pt. Date placeholder is dropped."""
    for ph in list(layout.placeholders):
        t = ph.placeholder_format.type
        if t == PP_PLACEHOLDER.DATE:
            ph._element.getparent().remove(ph._element)
        elif t == PP_PLACEHOLDER.FOOTER:
            if not show:
                ph._element.getparent().remove(ph._element)
                continue
            place(ph, LM, 6.98, 5, 0.28)
            style_ph(ph, 9, ink, MONO, anchor="ctr", cap=True, spc=140)
            ph.text_frame.text = "serfis.ai"
        elif t == PP_PLACEHOLDER.SLIDE_NUMBER:
            if not show:
                ph._element.getparent().remove(ph._element)
                continue
            place(ph, W_IN - LM - 1.5, 6.98, 1.5, 0.28)
            style_ph(ph, 9, ink, MONO, anchor="ctr", spc=140, algn="r")


def header_art(layout, ids, ink_line=C["line"], logo=True, hairline=True, eyebrow_y=1.22):
    art = []
    if logo:
        art.append(pic_xml(ids(), "Serfis logo", logo_rid(layout), LM, 0.32, 0.5, 0.5))
    if hairline:
        art.append(shape_xml(ids(), "Footer rule", LM, 6.9, CW, 0.0, geom="line", line=f'<a:ln w="9525">{solid(ink_line)}</a:ln>'))
    return art


def eyebrow(layout, ids, y=1.22, ink=C["text3"], x=LM, w=7.0):
    """Step tag: a short blue-to-lavender rule, then a mono caps label."""
    art = shape_xml(ids(), "Eyebrow rule", x, y + 0.135, 0.27, 0.02, fill=TAG_GRAD)
    layout.shapes._spTree.insert(2, X(art))
    add_placeholder(layout, placeholder_xml(ids(), "Eyebrow", 13, x + 0.4, y, w, 0.29,
                                            rpr(11, ink, MONO, spc=170, cap=True), "Eyebrow label"))


# =====================================================================================
# Layouts
# =====================================================================================
# remove placeholders' leftovers: dates everywhere; keep footer + slide numbers where wanted
# ---- 1. Title (hero) ----
ids = next_ids(200)
set_background(L_TITLE, f"{BG}/hero.jpg")
L_TITLE.name = "Serfis Title"
art = [pic_xml(ids(), "Serfis logo", logo_rid(L_TITLE), LM, 0.55, 0.78, 0.78),
       shape_xml(ids(), "Wordmark", LM + 0.92, 0.62, 3.5, 0.64, txbox=True, anchor="ctr",
                 text_xml=run_xml("Serfis", 30, C["text"], WORD, bold=True), insets=(0, 0, 0, 0)),
       shape_xml(ids(), "REC dot", W_IN - LM - 1.12, 0.84, 0.1, 0.1, geom="ellipse", fill=solid(C["rec"])),
       shape_xml(ids(), "REC label", W_IN - LM - 0.95, 0.74, 0.95, 0.3, txbox=True, anchor="ctr",
                 text_xml=run_xml("REC", 11, C["muted"], MONO, spc=180), insets=(0, 0, 0, 0))]
for which, (cx, cy) in dict(tl=(0.4, 0.4), tr=(W_IN - 0.4 - 0.32, 0.4), bl=(0.4, H_IN - 0.4 - 0.32), br=(W_IN - 0.4 - 0.32, H_IN - 0.4 - 0.32)).items():
    art.append(corner_xml(ids(), f"Frame {which}", cx, cy, 0.32, which))
decorate(L_TITLE, art)
place(ph_by_idx(L_TITLE, 0), LM, 2.55, 10.6, 2.3)
style_ph(ph_by_idx(L_TITLE, 0), 46, C["head"], DISPLAY, bold=True, anchor="b", cap=True, spc=-60)
place(ph_by_idx(L_TITLE, 1), LM, 5.1, 7.6, 0.9)
style_ph(ph_by_idx(L_TITLE, 1), 22, C["heliotrope"], BODY, bold=True, anchor="t", cap=False)
eyebrow(L_TITLE, ids, y=2.05)
finish_footers(L_TITLE, show=False)

# ---- 2. Section divider ----
ids = next_ids(200)
set_background(L_SECTION, f"{BG}/section.jpg")
L_SECTION.name = "Section Divider"
decorate(L_SECTION, header_art(L_SECTION, ids, hairline=True))
place(ph_by_idx(L_SECTION, 0), LM, 2.5, 10.8, 1.9)
style_ph(ph_by_idx(L_SECTION, 0), 44, C["head"], DISPLAY, bold=True, anchor="b", cap=True, spc=-50)
place(ph_by_idx(L_SECTION, 1), LM, 4.65, 7.8, 1.1)
style_ph(ph_by_idx(L_SECTION, 1), 20, C["text2"], BODY, anchor="t", cap=False)
eyebrow(L_SECTION, ids, y=2.0)
finish_footers(L_SECTION)


# ---- shared content-slide furniture ----
def content_layout(layout, name, bg=f"{BG}/content.jpg"):
    layout.name = name
    set_background(layout, bg)
    ids = next_ids(200)
    decorate(layout, header_art(layout, ids))
    eyebrow(layout, ids)
    t = layout.placeholders[0]
    place(t, LM, 1.58, CW, 1.15)
    style_ph(t, 30, C["head"], DISPLAY, bold=True, anchor="t", cap=True, spc=-30)
    finish_footers(layout)
    return ids


content_layout(L_CONTENT, "Title + Content")
b = ph_by_idx(L_CONTENT, 1)
place(b, LM, 3.0, CW, 3.6)
style_ph(b, 18, C["text2"], BODY, bullets="•", lnspc=112000)

content_layout(L_TWO, "Two Columns")
place(ph_by_idx(L_TWO, 1), LM, 3.0, 5.6, 3.6)
style_ph(ph_by_idx(L_TWO, 1), 17, C["text2"], BODY, bullets="•", lnspc=112000)
place(ph_by_idx(L_TWO, 2), LM + 6.03, 3.0, 5.6, 3.6)
style_ph(ph_by_idx(L_TWO, 2), 17, C["text2"], BODY, bullets="•", lnspc=112000)

ids = content_layout(L_COMPARE, "Comparison")
for idx, x in ((1, LM), (3, LM + 6.03)):
    p = ph_by_idx(L_COMPARE, idx)
    place(p, x, 2.95, 5.6, 0.45)
    style_ph(p, 16, C["heliotrope"], BODY, bold=True, anchor="b", cap=False)
for idx, x in ((2, LM), (4, LM + 6.03)):
    p = ph_by_idx(L_COMPARE, idx)
    place(p, x, 3.55, 5.6, 3.05)
    style_ph(p, 17, C["text2"], BODY, bullets="•", lnspc=112000)
decorate(L_COMPARE, [shape_xml(ids(), "Column rule L", LM, 3.5, 5.6, 0, geom="line", line=f'<a:ln w="9525">{solid(C["line"])}</a:ln>'),
                     shape_xml(ids(), "Column rule R", LM + 6.03, 3.5, 5.6, 0, geom="line", line=f'<a:ln w="9525">{solid(C["line"])}</a:ln>')])

content_layout(L_TITLEONLY, "Title Only")

# ---- Feature + Caption: title/text left, glass panel right ----
ids = content_layout(L_CAPTION, "Feature + Panel")
t = L_CAPTION.placeholders[0]
place(t, LM, 1.58, 4.9, 1.7)
style_ph(t, 26, C["head"], DISPLAY, bold=True, anchor="t", cap=True, spc=-30)
panel = shape_xml(ids(), "Glass panel", 6.35, 1.2, 6.13, 5.4, geom="roundRect", adj=5200,
                  fill=solid(C["glass"], 62), line=f'<a:ln w="9525">{solid(C["text"], 12)}</a:ln>')
decorate(L_CAPTION, [panel] + [corner_xml(ids(), f"HUD {w}", x, y, 0.22, w) for w, (x, y) in
                              dict(tl=(6.55, 1.4), tr=(12.28 - 0.0, 1.4), bl=(6.55, 6.18), br=(12.28, 6.18)).items()])
place(ph_by_idx(L_CAPTION, 1), 6.7, 1.6, 5.43, 4.6)
style_ph(ph_by_idx(L_CAPTION, 1), 17, C["text2"], BODY, bullets="•", lnspc=112000)
place(ph_by_idx(L_CAPTION, 2), LM, 3.45, 4.9, 3.1)
style_ph(ph_by_idx(L_CAPTION, 2), 16, C["text2"], BODY, anchor="t", cap=False, lnspc=118000)

# ---- Image + Caption: rounded picture left, title/text right ----
ids = content_layout(L_PICTURE, "Image + Caption")
pic = ph_by_idx(L_PICTURE, 1)
place(pic, LM, 1.25, 6.9, 5.35)
spPr = pic._element.spPr
for ch in list(spPr):
    if ch.tag == qn("a:prstGeom"):
        spPr.remove(ch)
spPr.append(parse_xml(f'<a:prstGeom {nsdecls("a")} prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 4000"/></a:avLst></a:prstGeom>'))
decorate(L_PICTURE, [corner_xml(ids(), f"HUD {w}", x, y, 0.26, w) for w, (x, y) in
                     dict(tl=(LM - 0.14, 1.11), tr=(LM + 6.9 - 0.12, 1.11), bl=(LM - 0.14, 6.46), br=(LM + 6.9 - 0.12, 6.46)).items()])
t = L_PICTURE.placeholders[0]
place(t, 8.3, 1.58, 4.18, 1.9)
style_ph(t, 24, C["head"], DISPLAY, bold=True, anchor="t", cap=True, spc=-30)
place(ph_by_idx(L_PICTURE, 2), 8.3, 3.65, 4.18, 2.9)
style_ph(ph_by_idx(L_PICTURE, 2), 16, C["text2"], BODY, anchor="t", cap=False, lnspc=118000)
# the eyebrow placeholder on this layout sits over the title column
for sp in L_PICTURE.shapes._spTree.iter(qn("p:sp")):
    nm = sp.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).get("name")
    if nm == "Eyebrow rule":
        off = sp.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:off")); off.set("x", str(E(8.3)))
    if nm == "Eyebrow":
        off = sp.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:off")); off.set("x", str(E(8.7)))
        ext = sp.find(qn("p:spPr")).find(qn("a:xfrm")).find(qn("a:ext")); ext.set("cx", str(E(3.7)))

# ---- Light (paper) ----
L_LIGHT.name = "Light: Title + Content"
strip_orient(L_LIGHT)
set_background(L_LIGHT, color=C["paper"])
ids = next_ids(200)
ov = L_LIGHT._element.find(qn("p:clrMapOvr"))
for ch in list(ov):
    ov.remove(ch)
ov.append(parse_xml(f'<a:overrideClrMapping {nsdecls("a")} bg1="lt1" tx1="dk1" bg2="lt2" tx2="dk2" accent1="accent1" accent2="accent2" '
                    f'accent3="accent3" accent4="accent4" accent5="accent5" accent6="accent6" hlink="hlink" folHlink="folHlink"/>'))
decorate(L_LIGHT, header_art(L_LIGHT, ids, ink_line="E6DEFA"))
eyebrow(L_LIGHT, ids, ink=C["on_paper"])
t = L_LIGHT.placeholders[0]
place(t, LM, 1.58, CW, 1.15)
style_ph(t, 30, C["on_paper"], DISPLAY, bold=True, anchor="t", cap=True, spc=-30)
b = ph_by_idx(L_LIGHT, 1)
place(b, LM, 3.0, CW, 3.6)
style_ph(b, 18, "3A2A66", BODY, bullets="•", lnspc=112000)
finish_footers(L_LIGHT, ink="8C7FB5")

# ---- Bright band (the site's one bright moment) ----
L_BAND.name = "Bright Band (closing)"
strip_orient(L_BAND)
set_background(L_BAND, f"{BG}/band.jpg")
ids = next_ids(200)
art = [shape_xml(ids(), "Footer rule", LM, 6.9, CW, 0.0, geom="line", line=f'<a:ln w="9525">{solid(C["ink"], 22)}</a:ln>')]
decorate(L_BAND, art)
t = L_BAND.placeholders[0]
place(t, LM, 2.2, 8.6, 2.4)
style_ph(t, 38, C["ink"], DISPLAY, bold=True, anchor="b", cap=True, spc=-40)
b = ph_by_idx(L_BAND, 1)
place(b, LM, 4.85, 7.2, 1.0)
style_ph(b, 20, C["ink"], BODY, anchor="t", cap=False, lnspc=118000)
eyebrow(L_BAND, ids, y=1.7, ink=C["ink"])
for sp in L_BAND.shapes._spTree.iter(qn("p:sp")):
    if sp.find(qn("p:nvSpPr")).find(qn("p:cNvPr")).get("name") == "Eyebrow rule":
        spPr = sp.find(qn("p:spPr"))
        spPr.remove(spPr.find(qn("a:gradFill")))
        spPr.find(qn("a:prstGeom")).addnext(parse_xml(f'<a:solidFill {nsdecls("a")}><a:srgbClr val="{C["ink"]}"/></a:solidFill>'))
finish_footers(L_BAND, ink=C["ink"])

# ---- Blank ----
L_BLANK.name = "Blank (dark)"
set_background(L_BLANK, f"{BG}/content.jpg")
finish_footers(L_BLANK, show=False)

# drop layouts we do not use (none left over), order the rest sensibly
order = [L_TITLE, L_SECTION, L_CONTENT, L_TWO, L_COMPARE, L_TITLEONLY, L_CAPTION, L_PICTURE, L_LIGHT, L_BAND, L_BLANK]
lst = master._element.find(qn("p:sldLayoutIdLst"))
by_part = {}
for el in list(lst):
    rid = el.get(qn("r:id"))
    by_part[master.part.related_part(rid).slide_layout.name] = el
for el in list(lst):
    lst.remove(el)
for l in order:
    lst.append(by_part[l.name])

# master placeholders: position to match the content layouts
for ph in master.placeholders:
    t = ph.placeholder_format.type
    if t == PP_PLACEHOLDER.TITLE:
        place(ph, LM, 1.58, CW, 1.15)
    elif t == PP_PLACEHOLDER.BODY:
        place(ph, LM, 3.0, CW, 3.6)
    elif t == PP_PLACEHOLDER.DATE:
        place(ph, W_IN - LM - 4.0, 6.98, 2.0, 0.28)
    elif t == PP_PLACEHOLDER.FOOTER:
        place(ph, LM, 6.98, 5, 0.28)
    elif t == PP_PLACEHOLDER.SLIDE_NUMBER:
        place(ph, W_IN - LM - 1.5, 6.98, 1.5, 0.28)


# =====================================================================================
# Sample slides
# =====================================================================================
def clone_footer_bits(slide):
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER):
            slide.shapes._spTree.append(copy.deepcopy(ph._element))


def new_slide(layout, title=None, eyebrow_text=None, footers=True):
    s = prs.slides.add_slide(layout)
    if title is not None and s.shapes.title is not None:
        s.shapes.title.text_frame.text = title
    if eyebrow_text is not None:
        s.placeholders[13].text_frame.text = eyebrow_text
    if footers:
        clone_footer_bits(s)
    return s


def set_paras(ph, items, bullets=True):
    tf = ph.text_frame
    tf.text = items[0]
    for t in items[1:]:
        tf.add_paragraph().text = t


def add_text(slide, x, y, w, h, runs, anchor="t", algn="l", lnspc=None):
    tb = slide.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    p = tf.paragraphs[0]
    p.alignment = {"l": PP_ALIGN.LEFT, "r": PP_ALIGN.RIGHT, "ctr": PP_ALIGN.CENTER}[algn]
    if lnspc:
        p.line_spacing = lnspc
    for r in runs:
        run = p.add_run()
        run.text = r["t"]
        f = run.font
        f.name = r["font"]; f.size = Pt(r["size"]); f.bold = r.get("bold", False)
        f.color.rgb = RGBColor.from_string(r["color"])
        rPr = run._r.get_or_add_rPr()
        if r.get("spc"):
            rPr.set("spc", str(r["spc"]))
        if r.get("cap"):
            rPr.set("cap", "all")
    return tb


def R(t, size, color, font=BODY, **kw):
    return dict(t=t, size=size, color=color, font=font, **kw)


def add_xml_shape(slide, xml):
    el = X(xml)
    slide.shapes._spTree.append(el)
    return el


def pill(slide, x, y, w, h, label, kind="primary", ids=[900]):
    """Pill button: gradient primary, ghost outline, or dark (for the bright band)."""
    ids[0] += 1
    if kind == "primary":
        fill, line, col = BTN_GRAD, f'<a:ln w="9525">{solid("FFFFFF", 20)}</a:ln>', "FFFFFF"
    elif kind == "ghost":
        fill, line, col = "<a:noFill/>", f'<a:ln w="12700">{solid(C["text"], 28)}</a:ln>', C["text"]
    else:
        fill, line, col = solid(C["ink"]), "<a:ln><a:noFill/></a:ln>", C["text"]
    add_xml_shape(slide, shape_xml(ids[0], f"Button {label}", x, y, w, h, geom="roundRect", adj=50000, fill=fill, line=line,
                                   text_xml=run_xml(label, 14, col, BODY, bold=True, spc=10), algn="ctr"))


def glass_card(slide, x, y, w, h, ids=[800]):
    ids[0] += 1
    add_xml_shape(slide, shape_xml(ids[0], "Glass card", x, y, w, h, geom="roundRect", adj=7500, fill=solid(C["glass"], 62),
                                   line=f'<a:ln w="9525">{solid(C["text"], 12)}</a:ln>'))


# ---- 1. Title ----
s = new_slide(L_TITLE, eyebrow_text="Eyebrow label  ·  Month 2026", footers=False)
tf = s.shapes.title.text_frame
tf.text = ""
p = tf.paragraphs[0]
for txt, col in (("Presentation title ", C["head"]), ("goes here", C["heliotrope"])):
    r = p.add_run(); r.text = txt; r.font.color.rgb = RGBColor.from_string(col)
s.placeholders[1].text_frame.text = "A one-line kicker that sets up the story."
pill(s, LM, 6.2, 2.3, 0.54, "Primary action", "primary")
pill(s, LM + 2.5, 6.2, 2.1, 0.54, "Secondary", "ghost")
s.notes_slide.notes_text_frame.text = (
    "Serfis template. Layouts: Title, Section Divider, Title + Content, Two Columns, Comparison, Title Only, "
    "Feature + Panel, Image + Caption, Light, Bright Band, Blank. Use Home > New Slide to pick one. "
    "Install the brand fonts first: Unbounded, DM Sans, DM Mono, Fredoka (all free on Google Fonts).")

# ---- 2. Agenda ----
s = new_slide(L_CONTENT, "Agenda", "Overview")
set_paras(s.placeholders[1], ["Where we are and why it matters", "What we built", "What we learned", "What comes next"])

# ---- 3. Section ----
s = new_slide(L_SECTION, "01 Section title", "Part one")
s.placeholders[1].text_frame.text = "One sentence that frames what this section covers."

# ---- 4. Numbered rows ----
s = new_slide(L_TITLEONLY, "How it works", "Process")
rows = [("01", "Step one", "One line that says what happens and why it matters."),
        ("02", "Step two", "Keep each description to a single sentence."),
        ("03", "Step three", "Rows are text boxes: copy, paste and edit."),
        ("04", "Step four", "Add or remove rows to fit the story.")]
y0, rh = 3.0, 0.9
for i, (n, ttl, desc) in enumerate(rows):
    y = y0 + i * rh
    add_xml_shape(s, shape_xml(700 + i, f"Row rule {n}", LM, y, CW, 0, geom="line", line=f'<a:ln w="9525">{solid(C["text"], 18)}</a:ln>'))
    add_text(s, LM, y + 0.22, 0.8, 0.4, [R(n, 13, C["heliotrope"], MONO, spc=140)])
    add_text(s, LM + 0.9, y + 0.17, 3.6, 0.5, [R(ttl, 20, C["text"], BODY, bold=True)])
    add_text(s, LM + 4.7, y + 0.2, CW - 4.7, 0.5, [R(desc, 16, C["text2"], BODY)])
add_xml_shape(s, shape_xml(710, "Row rule end", LM, y0 + 4 * rh, CW, 0, geom="line", line=f'<a:ln w="9525">{solid(C["text"], 18)}</a:ln>'))

# ---- 5. Two columns ----
s = new_slide(L_TWO, "Two columns", "Layout")
set_paras(s.placeholders[1], ["First column point", "Second point", "Third point"])
set_paras(s.placeholders[2], ["First column point", "Second point", "Third point"])

# ---- 6. Comparison ----
s = new_slide(L_COMPARE, "Before and after", "Comparison")
s.placeholders[1].text_frame.text = "Before"
set_paras(s.placeholders[2], ["Manual and slow", "Hard to measure", "Hard to repeat"])
s.placeholders[3].text_frame.text = "After"
set_paras(s.placeholders[4], ["Automated and fast", "Clear numbers", "Repeatable every time"])

# ---- 7. Metric cards ----
s = new_slide(L_TITLEONLY, "Three numbers that matter", "Results")
cw_, gap = (CW - 2 * 0.3) / 3, 0.3
for i, (num, lab, desc) in enumerate([("00", "Key metric", "A short line that explains the number."),
                                       ("00%", "Second metric", "A short line that explains the number."),
                                       ("0×", "Third metric", "A short line that explains the number.")]):
    x = LM + i * (cw_ + gap)
    glass_card(s, x, 3.0, cw_, 3.45)
    add_text(s, x + 0.35, 3.3, cw_ - 0.7, 0.35, [R(lab, 11, C["muted"], MONO, spc=170, cap=True)])
    add_text(s, x + 0.35, 3.85, cw_ - 0.7, 1.1, [R(num, 54, C["heliotrope"], DISPLAY, bold=True, spc=-60)], anchor="ctr")
    add_text(s, x + 0.35, 5.2, cw_ - 0.7, 1.0, [R(desc, 15, C["text2"], BODY)], lnspc=1.15)

# ---- 8. Feature + panel ----
s = new_slide(L_CAPTION, "Feature title", "Feature")
s.placeholders[2].text_frame.text = "Two or three sentences of supporting copy sit under the title on the left."
set_paras(s.placeholders[1], ["Point inside the glass panel", "Another point", "Drop a chart, table or screenshot here"])

# ---- 9. Chart ----
s = new_slide(L_TITLEONLY, "Chart title", "Data")
cd = CategoryChartData()
cd.categories = ["Q1", "Q2", "Q3", "Q4"]
cd.add_series("Series one", (4.2, 5.6, 7.1, 9.0))
cd.add_series("Series two", (3.1, 3.8, 4.4, 5.2))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, E(LM), E(2.9), E(CW), E(3.8), cd)
ch = gf.chart
ch.font.name, ch.font.size = BODY, Pt(12)
ch.font.color.rgb = RGBColor.from_string(C["muted"])
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.TOP
ch.legend.include_in_layout = False
ch.legend.font.color.rgb = RGBColor.from_string(C["muted"])
plot = ch.plots[0]
plot.gap_width = 70
plot.overlap = -10
for ser, col in zip(plot.series, (C["violet2"], C["heliotrope"])):
    ser.format.fill.solid(); ser.format.fill.fore_color.rgb = RGBColor.from_string(col)
va, ca = ch.value_axis, ch.category_axis
va.major_gridlines.format.line.color.rgb = RGBColor.from_string(C["line"])
va.format.line.fill.background()
ca.format.line.color.rgb = RGBColor.from_string(C["line"])
for ax in (va, ca):
    ax.tick_labels.font.color.rgb = RGBColor.from_string(C["muted"])
    ax.tick_labels.font.size = Pt(12)

cs = ch._chartSpace
cs.find(qn("c:chart")).addnext(parse_xml(
    '<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'))
pa = cs.find(qn("c:chart")).find(qn("c:plotArea"))
pa.append(parse_xml(
    '<c:spPr xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'))

# ---- 10. Image + caption ----
s = new_slide(L_PICTURE, "Image title", "Image")
s.placeholders[1].insert_picture(f"{BG}/sample_image.jpg")
s.placeholders[2].text_frame.text = "A short caption for the image on the left. Click the picture to replace it."

# ---- 11. Brand guide ----
s = new_slide(L_TITLEONLY, "Brand palette", "Brand guide")
sw = [("Ground", C["bg"], "06030F"), ("Night", C["bg2"], "0F032D"), ("Violet", C["violet"], "4E2BCC"), ("Violet 2", C["violet2"], "6A3FE6"),
      ("Lavender", C["lavender"], "905BF4"), ("Heliotrope", C["heliotrope"], "C084FC"), ("Blue", C["blue"], "3B5BFF"), ("Text", C["text"], "EFEFEF")]
sx = (CW - 7 * 0.2) / 8
for i, (nm, col, hx) in enumerate(sw):
    x = LM + i * (sx + 0.2)
    ids_ = 600 + i
    add_xml_shape(s, shape_xml(ids_, f"Swatch {nm}", x, 2.95, sx, 1.05, geom="roundRect", adj=14000, fill=solid(col),
                               line=f'<a:ln w="9525">{solid(C["text"], 22)}</a:ln>'))
    add_text(s, x, 4.1, sx, 0.25, [R(nm, 12, C["text"], BODY, bold=True)])
    add_text(s, x, 4.36, sx, 0.25, [R("#" + hx, 10, C["muted"], MONO, spc=60)])
add_xml_shape(s, shape_xml(650, "Button swatch", LM, 4.95, 2.6, 0.5, geom="roundRect", adj=50000, fill=BTN_GRAD,
                           text_xml=run_xml("Button gradient", 13, "FFFFFF", BODY, bold=True), algn="ctr"))
add_xml_shape(s, shape_xml(651, "Band swatch", LM + 2.8, 4.95, 2.6, 0.5, geom="roundRect", adj=50000,
                           fill=grad([(0, "4E2BCC"), (9, "9A66F8"), (58, "C58BFF"), (100, "D9B8FF")], 0),
                           text_xml=run_xml("Bright band", 13, C["ink"], BODY, bold=True), algn="ctr"))
add_text(s, LM, 5.75, 3.9, 0.5, [R("Unbounded", 22, C["head"], DISPLAY, bold=True, cap=True)])
add_text(s, LM, 6.2, 3.9, 0.3, [R("Headlines, uppercase, 700", 11, C["muted"], MONO, spc=60)])
add_text(s, LM + 4.1, 5.75, 3.9, 0.5, [R("DM Sans", 22, C["text"], BODY, bold=True)])
add_text(s, LM + 4.1, 6.2, 3.9, 0.3, [R("Body copy, 400 to 600", 11, C["muted"], MONO, spc=60)])
add_text(s, LM + 8.2, 5.75, 3.4, 0.5, [R("DM Mono", 20, C["heliotrope"], MONO, spc=60, cap=True)])
add_text(s, LM + 8.2, 6.2, 3.4, 0.3, [R("Labels and numbers, caps", 11, C["muted"], MONO, spc=60)])

# ---- 12. Light ----
s = new_slide(L_LIGHT, "Light variant", "For print and bright rooms")
set_paras(s.placeholders[1], ["Same layout system on a white ground", "Violet headings, deep-violet body copy", "Use it when a dark deck will not project well"])

# ---- 13. Bright band ----
s = new_slide(L_BAND, "Thank you", "Next step")
s.placeholders[1].text_frame.text = "name@serfis.ai  ·  serfis.ai"
pill(s, LM, 6.05, 2.3, 0.54, "Get in touch", "dark")

prs.core_properties.title = "Serfis presentation template"
prs.core_properties.author = "Serfis AI"
prs.core_properties.subject = "Serfis brand template"
prs.save(OUT)

# font-substitution hints (sans / monospace) so PowerPoint does not fall back to a serif when the brand fonts are missing
import os, zipfile
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.endswith(".xml") and item.filename.startswith("ppt/"):
            t = data.decode("utf8")
            t = re.sub(r'<a:(latin|cs) typeface="(Unbounded|DM Sans|Fredoka)"[^>]*/>', r'<a:\1 typeface="\2" pitchFamily="34" charset="0"/>', t)
            t = re.sub(r'<a:(latin|cs) typeface="(DM Mono)"[^>]*/>', r'<a:\1 typeface="\2" pitchFamily="49" charset="0"/>', t)
            data = t.encode("utf8")
        zout.writestr(item, data)
os.replace(tmp, OUT)
print("saved", OUT, "layouts:", [l.name for l in prs.slide_layouts], "slides:", len(prs.slides))
