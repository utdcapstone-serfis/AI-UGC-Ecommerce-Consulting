import copy, re
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Pt

SRC = str(Path(__file__).resolve().parents[1] / "Serfis_Template.pptx")
OUT = str(Path(__file__).resolve().parents[1] / "Serfis_Agents_Slide.pptx")

E = lambda i: int(round(i * 914400))
TEXT, TEXT2, MUTED, HELIO, GLASS = "EFEFEF", "B7B6B9", "A9A8AC", "C084FC", "10062C"
W, LM = 13.333, 0.85
CW = W - 2 * LM

prs = Presentation(SRC)
lst = prs.slides._sldIdLst
for el in list(lst):                       # drop the 13 sample slides
    prs.part.drop_rel(el.rId)
    lst.remove(el)

layout = next(l for l in prs.slide_layouts if l.name == "Title Only")
s = prs.slides.add_slide(layout)
s.shapes.title.text_frame.text = "Our nine agents"
s.placeholders[13].text_frame.text = "Lead-gen pipeline"
for ph in layout.placeholders:             # footer + slide number
    if ph.placeholder_format.type in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER):
        s.shapes._spTree.append(copy.deepcopy(ph._element))


def text(x, y, w, h, runs, anchor="t"):
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE}[anchor]
    p = tf.paragraphs[0]
    for t, size, color, font, bold, spc in runs:
        r = p.add_run(); r.text = t
        r.font.name, r.font.size, r.font.bold = font, Pt(size), bold
        r.font.color.rgb = RGBColor.from_string(color)
        if spc:
            r._r.get_or_add_rPr().set("spc", str(spc))
    return tb


agents = [
    ("Discovery", "Finds YC startups that fit our ideal customer"),
    ("Enrichment", "Adds size, industry, location and founders"),
    ("ICP guardrail", "Rejects poor fits before any spend"),
    ("Contact finder", "Decision-makers with published emails"),
    ("Scoring", "Rates customer fit from 1 to 10"),
    ("Research", "Finds talking points and the best ad type"),
    ("Outreach writer", "Drafts the first email in our voice"),
    ("Reply writer", "Drafts replies and honors opt-outs"),
    ("Send agent", "Sends approved emails only, capped daily"),
]
gap, rh, y0 = 0.25, 1.1, 2.75
cw = (CW - 2 * gap) / 3
for i, (name, desc) in enumerate(agents):
    r_, c_ = divmod(i, 3)
    x, y = LM + c_ * (cw + gap), y0 + r_ * (rh + 0.14)
    s.shapes._spTree.append(parse_xml(
        f'<p:sp {nsdecls("a", "p")}><p:nvSpPr><p:cNvPr id="{300 + i}" name="Agent card {i + 1}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        f'<p:spPr><a:xfrm><a:off x="{E(x)}" y="{E(y)}"/><a:ext cx="{E(cw)}" cy="{E(rh)}"/></a:xfrm>'
        f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val 12000"/></a:avLst></a:prstGeom>'
        f'<a:solidFill><a:srgbClr val="{GLASS}"><a:alpha val="62000"/></a:srgbClr></a:solidFill>'
        f'<a:ln w="9525"><a:solidFill><a:srgbClr val="{TEXT}"><a:alpha val="12000"/></a:srgbClr></a:solidFill></a:ln></p:spPr></p:sp>'))
    text(x + 0.25, y + 0.2, 0.5, 0.3, [(f"{i + 1:02d}", 11, HELIO, "DM Mono", False, 140)])
    text(x + 0.75, y + 0.17, cw - 1.0, 0.36, [(name, 17, TEXT, "DM Sans", True, 0)])
    text(x + 0.25, y + 0.62, cw - 0.5, 0.42, [(desc, 13, TEXT2, "DM Sans", False, 0)])

text(LM, 6.5, CW, 0.28, [("Nothing is emailed until a person approves it", 10, MUTED, "DM Mono", False, 120)])

prs.save(OUT)

# same font-substitution hints as the template (sans / mono fallbacks)
import os, zipfile
tmp = OUT + ".tmp"
with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename.endswith(".xml") and item.filename.startswith("ppt/"):
            t = data.decode("utf8")
            t = re.sub(r'<a:(latin|cs) typeface="(Unbounded|DM Sans|Fredoka)"[^>]*/>', r'<a:\1 typeface="\2" pitchFamily="34" charset="0"/>', t)
            t = re.sub(r'<a:(latin|cs) typeface="(DM Mono)"[^>]*/>', r'<a:\1 typeface="\2" pitchFamily="49" charset="0"/>', t)
            data = t.encode("utf8")
        zout.writestr(item, data)
os.replace(tmp, OUT)
print("saved", OUT, len(Presentation(OUT).slides), "slide")
