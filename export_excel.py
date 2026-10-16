"""Exports all leads to a formatted Excel workbook for human review/approval."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
import storage

COLUMNS = [
    ("company_name", "Company", 22),
    ("website", "Website", 26),
    ("industry", "Industry", 18),
    ("company_size", "Company Size", 14),
    ("location", "Location", 18),
    ("what_they_sell", "What They Sell", 30),
    ("funding_notes", "Funding Notes", 22),
    ("priority", "Priority", 9),
    ("category", "Category", 22),
    ("market", "Market", 14),
    ("serfis_fit", "Serfis Fit", 14),
    ("icp_fit_score", "ICP Fit Score (1-10)", 10),
    ("why_it_fits", "Why It Fits", 34),
    ("contact_name", "Contact Name", 18),
    ("contact_title", "Contact Title", 20),
    ("contact_email", "Contact Email", 24),
    ("email_status", "Email Status", 14),
    ("contact_check", "How Contact Was Verified", 46),
    ("contact_profile_url", "Contact/Profile URL", 26),
    ("research_notes", "Research Notes", 40),
    ("research_sources", "Research Sources", 30),
    ("content_angle", "Suggested Ad / Content Angle", 36),
    ("best_ad_type", "Best Ad Type", 22),
    ("founders", "Founders (YC)", 36),
    ("yc_batch", "YC Batch", 12),
    ("decided_at", "Approval Decision Time (UTC)", 20),
    ("overrides", "Guardrail Flags Acknowledged", 18),
    ("yc_profile_url", "YC Profile", 34),
    ("outreach_angle", "Recommended Outreach Angle", 34),
    ("draft_email_subject", "Draft Email Subject", 26),
    ("draft_email_body", "Draft Email Body", 50),
    ("lead_status", "Lead Status", 14),
    ("approval_status", "Approval Status", 16),
]


def export_to_excel(path=None) -> str:
    path = path or config.OUTPUT_XLSX
    leads = storage.get_all_leads()

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    for col_idx, (_, header, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    for row_idx, lead in enumerate(leads, start=2):
        data = lead.as_dict()
        for col_idx, (field, _, _) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=data.get(field))
            cell.alignment = wrap

    wb.save(path)
    return str(path)


if __name__ == "__main__":
    out_path = export_to_excel()
    print(f"Exported {len(storage.get_all_leads())} leads to {out_path}")
